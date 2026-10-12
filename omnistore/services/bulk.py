"""Sequential catalog ingestion from Excel, CSV, DOCX, and Google Sheets."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any, Callable
from urllib.parse import parse_qs, urlparse

import pandas as pd

from ..models import UniversalProduct

MAX_BULK_ROWS = 1000
SHEETS_RE = re.compile(r"docs\.google\.com/spreadsheets/d/([A-Za-z0-9_-]+)")


def _records_from_google_sheet(url: str, google_credentials: dict[str, Any]) -> list[dict[str, Any]]:
    match = SHEETS_RE.search(url)
    if not match:
        raise ValueError("That is not a valid Google Sheets URL")
    spreadsheet_id = match.group(1)
    query = parse_qs(urlparse(url).query)
    gid = query.get("gid", ["0"])[0]
    api_key = str(google_credentials.get("api_key") or "").strip()
    service_info = google_credentials.get("service_account_json")
    kwargs: dict[str, Any] = {"cache_discovery": False}
    if service_info:
        from google.oauth2.service_account import Credentials

        if isinstance(service_info, str):
            import json
            service_info = json.loads(service_info)
        creds = Credentials.from_service_account_info(
            service_info,
            scopes=["https://www.googleapis.com/auth/spreadsheets.readonly"],
        )
        kwargs["credentials"] = creds
    elif api_key:
        kwargs["developerKey"] = api_key
    else:
        raise ValueError("Google Sheets access needs a Google API key (public sheet) or service account JSON")

    from googleapiclient.discovery import build

    service = build("sheets", "v4", **kwargs)
    # The API accepts a sheet-qualified A1 range. A numeric gid alone is not an A1 title,
    # so use spreadsheet metadata to map the gid to a safe, escaped tab title.
    metadata = service.spreadsheets().get(spreadsheetId=spreadsheet_id).execute()
    sheet_title = next((s["properties"]["title"] for s in metadata.get("sheets", [])
                        if str(s["properties"].get("sheetId")) == str(gid)), None)
    if not sheet_title:
        raise ValueError(f"No sheet tab with gid={gid} was found")
    escaped_title = sheet_title.replace("'", "''")
    values = service.spreadsheets().values().get(
        spreadsheetId=spreadsheet_id, range=f"'{escaped_title}'!A:ZZ"
    ).execute().get("values", [])
    if not values:
        return []
    headers = [str(value).strip().lower().replace(" ", "_") for value in values[0]]
    return [dict(zip(headers, row)) for row in values[1:] if any(str(cell).strip() for cell in row)]


def load_rows(file_source: str | Path, google_credentials: dict[str, Any] | None = None) -> list[dict[str, Any]]:
    source = str(file_source).strip()
    if SHEETS_RE.search(source):
        rows = _records_from_google_sheet(source, google_credentials or {})
    else:
        path = Path(source).expanduser().resolve()
        if not path.is_file():
            raise FileNotFoundError(f"Bulk source file was not found: {path}")
        suffix = path.suffix.lower()
        if suffix == ".xlsx":
            frame = pd.read_excel(path, engine="openpyxl", dtype=str).fillna("")
            frame.columns = [str(column).strip().lower().replace(" ", "_") for column in frame.columns]
            rows = frame.to_dict(orient="records")
        elif suffix == ".csv":
            frame = pd.read_csv(path, dtype=str, keep_default_na=False)
            frame.columns = [str(column).strip().lower().replace(" ", "_") for column in frame.columns]
            rows = frame.to_dict(orient="records")
        elif suffix == ".docx":
            from docx import Document

            document = Document(path)
            if not document.tables:
                raise ValueError("DOCX imports need a table with column headers in the first row")
            rows = []
            for table in document.tables:
                table_rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
                if len(table_rows) < 2:
                    continue
                headers = [value.lower().replace(" ", "_") for value in table_rows[0]]
                rows.extend(dict(zip(headers, row)) for row in table_rows[1:] if any(row))
        else:
            raise ValueError("Supported bulk files are .xlsx, .csv, .docx, or a Google Sheets URL")
    if not rows:
        raise ValueError("The bulk source did not contain any product rows")
    if len(rows) > MAX_BULK_ROWS:
        raise ValueError(f"The source has {len(rows)} rows; the limit is {MAX_BULK_ROWS} per run")
    return rows


def process_bulk_file(
    file_source: str | Path,
    target_platforms: list[str],
    *,
    google_credentials: dict[str, Any],
    resolve_adapter: Callable[[str], Any],
    enrich_product: Callable[[dict[str, Any]], UniversalProduct],
    progress: Callable[[int, int, str], None] | None = None,
) -> dict[str, Any]:
    """Read once, AI-normalize once per row, then sequentially fan out per channel."""
    rows = load_rows(file_source, google_credentials)
    targets = list(dict.fromkeys(str(name).lower() for name in target_platforms if str(name).strip()))
    if not targets:
        raise ValueError("Choose one or more target platforms for this import")
    results: list[dict[str, Any]] = []
    total = len(rows) * len(targets)
    completed = 0
    for row_index, row in enumerate(rows, start=1):
        # The same canonical object is sent to each adapter; channel mappings cannot drift.
        product = enrich_product(row)
        for platform in targets:
            try:
                adapter = resolve_adapter(platform)
                result = adapter.create_product(product)
                results.append({"row": row_index, "platform": platform, "ok": True,
                                "product": product.title, "result": result})
                status = f"Row {row_index}/{len(rows)} listed on {platform}: {product.title}"
            except Exception as exc:
                results.append({"row": row_index, "platform": platform, "ok": False,
                                "product": product.title, "error": str(exc)})
                status = f"Row {row_index}/{len(rows)} failed on {platform}: {exc}"
            completed += 1
            if progress:
                progress(completed, total, status)
    succeeded = sum(1 for item in results if item["ok"])
    return {"rows": len(rows), "platforms": targets, "succeeded": succeeded,
            "failed": len(results) - succeeded, "results": results}
